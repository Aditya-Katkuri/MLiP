#!/usr/bin/env python3
"""Build the 5-slide, 3-minute pitch deck for Sidewalk Scanner (16:9; the speaker notes are the word-for-word script).

Inputs come from build_assets.py and build_charts.py, in reports/presentation/build/.
usage: /data/tools/shotenv/bin/python reports/presentation/build_deck.py && bash reports/presentation/render.sh
"""
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
B, IMG = HERE / "build", HERE / "build" / "img"
OUT = HERE / "sidewalk-scanner-pitch.pptx"

INK, BODY, MUTED, RULE, SOFT, WHITE = "111827", "374151", "6B7280", "D1D5DB", "F3F4F6", "FFFFFF"
TEAL, TEAL_SOFT, SLATE = "0F766E", "E6F2F0", "475569"
TIER = {1: "DC2626", 2: "F59E0B", 3: "7C8A9A"}
FONT = "Calibri"
ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

# Pittsburgh residents, U.S. Census ACS 2024 1-year, City of Pittsburgh: B18105 (ambulatory difficulty, age 5+) and B01001 (age)
AMBULATORY_DIFFICULTY, AGE_65_PLUS, UNDER_5 = 19_194, 54_826, 14_733
# CMAP, "Measuring sidewalk accessibility" (21 May 2026): $1,600/mile field surveys, $64/mile with Project Sidewalk;
# municipalities C and D, (54 + 130) sidewalk miles / (565 + 692) collection hours x 40 h = 5.86 miles per staff-week
COST_FIELD, COST_STREET_VIEW, MILES_PER_STAFF_WEEK = 1_600, 64, 5.86
# /data/eda/baseline/run.log: CLIP feature extraction ran at 37 photos/s on the CPU VM, so 21,071 photos took ~9.5 minutes

# Speaker notes: what to say, point by point, for a 3-minute talk (~430 words at ~155 words per minute)
NOTES = {
    1: ["Hi, we're Team YOLO, and this is Sidewalk Scanner.",
        "Sidewalk barriers hit nineteen thousand Pittsburghers who struggle to walk, fifty-five thousand seniors, fifteen thousand kids "
        "in strollers, and their caregivers.",
        "One missing curb ramp, or a pole in the sidewalk, and a wheelchair can't get through.",
        "The City needs these barriers for its new ADA plan, but finding them is slow.",
        "Walking surveys cost sixteen hundred dollars a mile.",
        "Checking Street View by hand covers about six miles a week per person.",
        "After six years, only seventeen percent of Pittsburgh is checked. That's the teal on the map.",
        "A model is far faster: ours checked twenty-one thousand photos in under ten minutes."],
    2: ["Here's how it works.",
        "A street photo goes in, and the model says: missing curb ramp, obstacle, surface problem, or no barrier.",
        "Every photo has GPS, so we know exactly where each barrier is.",
        "Each barrier then gets a tier. Tier one is near hospitals, schools and busy bus stops.",
        "So instead of waiting years, the City gets one citywide list, and planners check Tier one first."],
    3: ["We don't have to label anything to start.",
        "Project Sidewalk has published about twenty-one thousand street photos from nineteen cities, each voted on by the crowd.",
        "We train on eighteen cities and test on Pittsburgh's five hundred fifty-nine photos.",
        "Pittsburgh's Project Sidewalk data tells us where each photo is.",
        "City open data on hospitals, schools and bus stops gives us the tiers.",
        "We also have RampNet, about two hundred fourteen thousand panoramas, but it only tells us whether a curb ramp is there."],
    4: ["We score it only on Pittsburgh, with two questions for each tier.",
        "Found: of the real barriers, how many did we catch?",
        "Right: of our flags, how many are real?",
        "For the truth, we use five hundred fifty-nine Pittsburgh photos the model never sees in training.",
        "We hand-label three hundred Mapillary photos, and we visit sampled spots to check them ourselves.",
        "Our targets: at least eighty percent of Tier one flags are real, and we catch more obstacles and surface problems than "
        "human labellers."],
    5: ["Our biggest risk is that a model trained in other cities won't work here.",
        "RampNet comes from New York, Portland and Bend, but Pittsburgh has hills, over a thousand public stairways, and brick sidewalks.",
        "That's why we test only on Pittsburgh and check spots in person.",
        "We also don't know yet how many unaudited streets have usable photos.",
        "And labels are sparse: a barrier is labelled once, but shows up in many photos.",
        "We'd love help with imagery, a DOMI planner, disability advocates, and GPU time. Thank you!"],
}


def rgb(h):
    return RGBColor.from_string(h)


def text(s, x, y, w, h, paras, size=16, color=INK, bold=False, align="l", anchor="t", spacing=None, after=0):
    """paras: a str, or a list of paragraphs; a paragraph is a str or a list of runs; a run is a str or (str, options)."""
    shp = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shp.text_frame
    tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[anchor]
    for i, para in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        if spacing:
            p.line_spacing = spacing
        if after:
            p.space_after = Pt(after)
        for run in [para] if isinstance(para, (str, tuple)) else para:
            t, o = (run, {}) if isinstance(run, str) else run
            r = p.add_run()
            r.text = t
            f = r.font
            f.name, f.size, f.bold, f.italic = FONT, Pt(o.get("size", size)), o.get("bold", bold), o.get("italic", False)
            f.color.rgb = rgb(o.get("color", color))
            if o.get("spc"):
                f._element.set("spc", str(o["spc"]))
    return shp


def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, radius=0, shape=None, alpha=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    shp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = shp._element.spPr.find(qn("a:solidFill"))[0]
            clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))}))
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)
    return shp


def chip(s, x, y, w, h, label, fill, color, size=12, bold=True, line=None, radius=None):
    rect(s, x, y, w, h, fill=fill, line=line, radius=h / 2 if radius is None else radius)
    text(s, x, y, w, h, label, size=size, bold=bold, color=color, align="c", anchor="m")


def picture(s, path, x, y, w, h=None, border=None):
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h) if h else None)
    if border:
        pic.line.color.rgb = rgb(border)
        pic.line.width = Pt(0.75)
    return pic


def fit(path, w, h):
    iw, ih = Image.open(path).size
    k = min(w / iw, h / ih)
    return iw * k, ih * k


def arrow(s, x1, y1, x2, y2, color=MUTED, lw=2.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    style = c._element.find(qn("p:style"))
    if style is not None:
        c._element.remove(style)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def label(s, x, y, w, t, color=MUTED, size=11):
    text(s, x, y, w, 0.28, [[(t.upper(), {"spc": 80})]], size=size, bold=True, color=color)


def slide(prs, n, kicker, title, notes, right="Sidewalk Scanner  ·  Team YOLO"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.6, 0.36, 7.5, 0.3, [[(f"{n:02d}   ", {"color": TEAL, "spc": 80}), (kicker.upper(), {"color": MUTED, "spc": 80})]], size=12, bold=True)
    text(s, 5.2, 0.36, 7.53, 0.3, right, size=11, color=MUTED, align="r")
    text(s, 0.6, 0.66, 12.13, 0.75, title, size=28, bold=True)
    text(s, 12.08, 7.16, 0.65, 0.22, f"{n} / 5", size=9, color=MUTED, align="r")
    words = sum(len(line.split()) for line in notes)
    seconds = round(words / 155 * 60 / 5) * 5
    s.notes_slide.notes_text_frame.text = "\n".join([f"(about {seconds} seconds)"] + [f"•  {line}" for line in notes])
    return s


def main():
    A = json.load(open(B / "assets.json"))
    M = json.load(open(B / "metrics.json"))
    W = json.load(open(B / "worklist.json"))
    km_share = {int(k): v for k, v in A["km_share_by_tier"].items()}
    oc, pc = M["class_counts_other"], M["class_counts_pgh"]
    n_photos = M["n_other"] + M["n_test_pgh"]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # standard 16:9 widescreen
    prs.part._element.find(qn("p:sldSz")).attrib.pop("type", None)  # the default template still says "screen4x3"

    # ---------------------------------------------------------------- 1. problem: finding barriers by hand is slow
    s = slide(prs, 1, "The problem  ·  who has it", "Pittsburgh still finds sidewalk barriers by hand, and it's slow",
              NOTES[1],
              right="Team YOLO  ·  Aditya Katkuri  ·  Prateik Sinha  ·  Giridhar Vadhul  ·  JS Durand")
    mw, mh = fit(B / "map_coverage.png", 6.5, 5.55)
    mx, my = 0.35, 1.6
    picture(s, B / "map_coverage.png", mx, my, mw, mh)
    text(s, mx + 0.1, my + 0.05, 1.5, 1.3, [[(f"{A['audited_share']:.0%}", {"size": 34, "bold": True, "color": TEAL})],
                                            "of streets checked", "by hand in 6 years"], size=12.5, color=BODY)

    X, WR = 7.05, 5.68
    label(s, X, 1.55, WR, "How barriers are found today")
    ways = [("On foot", "A surveyor walks every block", f"${COST_FIELD:,} / mile", "field survey", SOFT, INK),
            ("Street View, by hand", "People click through panoramas", f"${COST_STREET_VIEW} / mile",
             f"≈ {MILES_PER_STAFF_WEEK:.0f} miles per staff-week", SOFT, INK),
            ("Automated sweep", "A model checks every photo", f"{n_photos:,} photos", "checked in under 10 minutes", TEAL_SOFT, TEAL)]
    for k, (name, how, value, sub, fill, vcolor) in enumerate(ways):
        yy = 1.9 + k * 0.7
        rect(s, X, yy, WR, 0.62, fill=fill, radius=0.08)
        text(s, X + 0.2, yy + 0.07, 3.2, 0.28, name, size=15, bold=True, color=TEAL if k == 2 else INK)
        text(s, X + 0.2, yy + 0.35, 3.2, 0.24, how, size=11.5, color=MUTED)
        text(s, X + 3.2, yy + 0.05, WR - 3.4, 0.32, value, size=18, bold=True, color=vcolor, align="r")
        text(s, X + 3.2, yy + 0.36, WR - 3.4, 0.22, sub, size=10.5, color=MUTED, align="r")

    PW = 1.62
    pgap = (WR - 3 * PW) / 2
    for i, (img, cap) in enumerate([("ncr_1779", "Missing curb ramp"), ("obs_10574", "Obstacle"), ("sfc_14216", "Surface problem")]):
        xx = X + i * (PW + pgap)
        picture(s, IMG / f"{img}.jpg", xx, 4.08, PW, PW)
        text(s, xx - 0.1, 5.73, PW + 0.2, 0.28, cap, size=13, bold=True, align="c")

    label(s, X, 6.1, WR, "Who gets stuck, and their caregivers")
    TW = (WR - 0.2) / 3
    for k, (num, who) in enumerate([(AMBULATORY_DIFFICULTY, "struggle to walk"), (AGE_65_PLUS, "seniors, 65 and over"),
                                    (UNDER_5, "kids under 5, in strollers")]):
        xx = X + k * (TW + 0.1)
        rect(s, xx, 6.42, TW, 0.72, fill=SOFT, radius=0.08)
        text(s, xx + 0.14, 6.46, TW - 0.2, 0.32, f"{num:,}", size=18, bold=True, color=TEAL)
        text(s, xx + 0.14, 6.8, TW - 0.2, 0.3, who, size=11, color=BODY)

    # ---------------------------------------------------------------- 2. system
    s = slide(prs, 2, "What we build  ·  what changes", "We turn street photos into a barrier list, sorted by tier",
              NOTES[2])
    CW, GAP, Y0, CH = 2.13, 0.37, 1.98, 2.1
    xs = [0.6 + i * (CW + GAP) for i in range(5)]
    for i, step in enumerate(["Photo", "Classify", "Locate", "Tier", "Worklist"]):
        text(s, xs[i], 1.58, CW, 0.34, [[(f"{i + 1}  ", {"color": TEAL}), step]], size=16, bold=True)
        if i < 4:
            arrow(s, xs[i] + CW + 0.07, Y0 + CH / 2, xs[i + 1] - 0.07, Y0 + CH / 2)
    picture(s, IMG / "ncr_23869.jpg", xs[0] + 0.015, Y0, 2.1, 2.1)
    text(s, xs[0], 4.16, CW, 0.5, "Street-level photo of a sidewalk", size=11, color=MUTED)

    rect(s, xs[1], Y0, CW, CH, fill=SOFT, radius=0.1)
    for j, c in enumerate(["Missing curb ramp", "Obstacle", "Surface problem", "No barrier"]):
        hit = j == 0
        chip(s, xs[1] + 0.13, Y0 + 0.16 + j * 0.48, CW - 0.26, 0.38, c,
             fill=TEAL if hit else WHITE, color=WHITE if hit else BODY, size=12.5, bold=hit, line=None if hit else RULE)
    text(s, xs[1], 4.16, CW, 0.5, "Image classifier picks one of four classes", size=11, color=MUTED)

    rect(s, xs[2], Y0, CW, CH, fill=SOFT, radius=0.1)
    rect(s, xs[2] + CW / 2 - 0.23, Y0 + 0.2, 0.46, 0.46, fill=TEAL, shape=MSO_SHAPE.OVAL)
    rect(s, xs[2] + CW / 2 - 0.08, Y0 + 0.35, 0.16, 0.16, fill=WHITE, shape=MSO_SHAPE.OVAL)
    text(s, xs[2] + 0.08, Y0 + 0.8, CW - 0.16, 1.25,
         [[("40.4463° N, 79.9438° W", {"bold": True, "color": INK})], "Squirrel Hill North",
          [("136 m", {"bold": True, "color": INK}), " to a preschool"], [("143 m", {"bold": True, "color": INK}), " to a bus stop"]],
         size=11.5, color=BODY, align="c", after=2)
    text(s, xs[2], 4.16, CW, 0.5, "Every photo carries its coordinates", size=11, color=MUTED)

    picture(s, B / "map_tiers_zoom.png", xs[3] + 0.015, Y0, 2.1, 2.1, border=RULE)
    chip(s, xs[3] + 0.12, Y0 + 0.12, 0.8, 0.3, "Tier 1", fill=TIER[1], color=WHITE, size=12)
    for k in (1, 2, 3):
        lx = xs[3] + (k - 1) * 0.71
        rect(s, lx, 4.23, 0.13, 0.13, fill=TIER[k])
        text(s, lx + 0.18, 4.15, 0.55, 0.28, f"Tier {k}", size=11, color=MUTED)
    text(s, xs[3], 4.42, CW, 0.28, "■ school, clinic, care   ○ busy stop", size=10, color=MUTED)

    rect(s, xs[4], Y0, CW, CH, fill=WHITE, line=RULE, radius=0.1)
    label(s, xs[4] + 0.14, Y0 + 0.1, CW - 0.28, "City worklist", size=10)
    for j, r in enumerate(W[:5]):
        ry = Y0 + 0.42 + j * 0.33
        rect(s, xs[4] + 0.14, ry + 0.07, 0.14, 0.14, fill=TIER[r["tier"]], shape=MSO_SHAPE.OVAL)
        text(s, xs[4] + 0.36, ry, CW - 0.44, 0.3, [[(r["type"], {"bold": True, "color": INK}), ("  " + r["hood"], {"color": MUTED, "size": 10})]], size=11.5)
    text(s, xs[4], 4.16, CW, 0.5, "Planners check Tier 1 first, then schedule repairs", size=11, color=MUTED)

    PY, PH = 4.95, 1.9
    rect(s, 0.6, PY, 5.75, PH, fill=SOFT, radius=0.1)
    label(s, 0.88, PY + 0.2, 5.2, "Today")
    text(s, 0.88, PY + 0.56, 5.2, 1.25, ["•  People walk streets or click through Street View", f"•  About {MILES_PER_STAFF_WEEK:.0f} miles per staff-week",
                                          "•  17% of Pittsburgh checked in 6 years"], size=16, color=BODY, after=4)
    arrow(s, 6.45, PY + PH / 2, 6.88, PY + PH / 2, color=TEAL, lw=2.5)
    rect(s, 6.98, PY, 5.75, PH, fill=TEAL_SOFT, radius=0.1)
    label(s, 7.26, PY + 0.2, 5.2, "With Sidewalk Scanner", color=TEAL)
    text(s, 7.26, PY + 0.56, 5.3, 1.25, ["•  A model sweeps every street photo", "•  Re-run whenever new imagery arrives",
                                          [f"•  Planners check Tier 1 first: {km_share[1]:.0%} of street km"]], size=16, color=INK, after=4)

    # ---------------------------------------------------------------- 3. data
    s = slide(prs, 3, "Data", "We learn from 18 cities and test on Pittsburgh",
              NOTES[3])
    IW, IG = 1.6, 0.15
    cols = [("Missing curb ramp", "teaneck_11267", "Teaneck", "missing_curb_ramp", "ncr_9511"),
            ("Obstacle", "sea_169860", "Seattle", "obstacle", "obs_16325"),
            ("Surface problem", "chicago_121128", "Chicago", "surface_problem", "sfc_5178"),
            ("No barrier*", "taipei_44530", "Taipei", "no_barrier", "rej_sfc_15791")]
    rows_y = (2.38, 4.62)
    label(s, 0.6, 2.02, 6.85, "18 other cities  ·  training")
    label(s, 0.6, 4.26, 6.85, "Pittsburgh  ·  held-out test", color=TEAL)
    for i, (name, other, city, key, pgh_img) in enumerate(cols):
        xx = 0.6 + i * (IW + IG)
        text(s, xx, 1.6, IW, 0.32, name, size=14, bold=True)
        for y, img, tag in [(rows_y[0], IMG / f"{other}.jpg", f"{city}  ·  {oc[key]:,}"), (rows_y[1], IMG / f"{pgh_img}.jpg", f"Pittsburgh  ·  {pc[key]:,}")]:
            picture(s, img, xx, y, IW, IW)
            rect(s, xx, y + IW - 0.28, IW, 0.28, fill=INK, alpha=0.72)
            text(s, xx + 0.08, y + IW - 0.28, IW - 0.12, 0.28, tag, size=10.5, bold=True, color=WHITE, anchor="m")
    text(s, 0.6, 6.42, 6.9, 0.4, [[("Label = crowd verdict. ", {"bold": True, "color": INK}), "Confirmed → that barrier.  *Rejected → “no barrier”."]],
         size=13, color=BODY)

    X, WR = 7.85, 4.88
    label(s, X, 1.65, WR, "Split by city, not at random")
    train_w = WR - 0.42
    rect(s, X, 2.02, train_w, 0.5, fill=SLATE, radius=0.06)
    text(s, X + 0.16, 2.02, train_w - 0.3, 0.5, [[("Train  ", {"bold": True}), f"18 cities  ·  {M['n_other']:,} photos"]], size=13, color=WHITE, anchor="m")
    rect(s, X + train_w + 0.1, 2.02, 0.32, 0.5, fill=TEAL, radius=0.06)
    text(s, X, 2.57, WR, 0.3, [[("Test  ", {"bold": True}), f"Pittsburgh  ·  {M['n_test_pgh']} photos"]], size=12, color=TEAL, align="r")
    label(s, X, 3.12, WR, "Where each piece comes from")
    sources = [("Project Sidewalk photos", f"{n_photos:,} crowd-labelled Street View crops, 19 cities"),
               ("Project Sidewalk Pittsburgh API", "25,690 labels with GPS: where each test photo is"),
               ("WPRDC + PRT open data", "Hospitals, schools, clinics, 2,696 bus stops: the tiers"),
               ("RampNet dataset", ["214,376 street panoramas (New York, Portland, Bend)", "Only detects whether a curb ramp is there"])]
    for k, (title, detail) in enumerate(sources):
        yy = 3.5 + k * 0.84
        chip(s, X, yy + 0.02, 0.38, 0.38, str(k + 1), fill=TEAL, color=WHITE, size=13)
        text(s, X + 0.55, yy, WR - 0.55, 0.3, title, size=14.5, bold=True)
        text(s, X + 0.55, yy + 0.31, WR - 0.55, 0.5, detail, size=12, color=BODY)

    # ---------------------------------------------------------------- 4. evaluation
    s = slide(prs, 4, "Evaluation", "Scored on Pittsburgh only: per barrier type, per tier",
              NOTES[4])
    LW = 7.75
    label(s, 0.6, 1.65, LW, "Two numbers, per tier")
    QW = (LW - 0.25) / 2
    for k, (big, body, term) in enumerate([("Found", "Of the real barriers in Tier 1, what share did we flag?", "recall"),
                                           ("Right", "Of our Tier 1 flags, what share are real barriers?", "precision")]):
        xx = 0.6 + k * (QW + 0.25)
        rect(s, xx, 2.02, QW, 1.62, fill=SOFT, radius=0.1)
        text(s, xx + 0.25, 2.16, 2.2, 0.55, big, size=28, bold=True, color=TEAL)
        text(s, xx + QW - 1.85, 2.3, 1.6, 0.35, f"({term})", size=12, color=MUTED, align="r")
        text(s, xx + 0.25, 2.8, QW - 0.5, 0.75, body, size=16, color=BODY)

    label(s, 0.6, 3.98, LW, "How we check the truth")
    GW = (LW - 0.5) / 3
    for k, (big, body) in enumerate([(f"{M['n_test_pgh']}", "held-out Pittsburgh photos with crowd votes"),
                                     ("300", "Mapillary photos we hand-label"),
                                     ("In person", "We go to sampled spots and check them ourselves")]):
        xx = 0.6 + k * (GW + 0.25)
        rect(s, xx, 4.35, GW, 2.47, fill=WHITE, line=RULE, radius=0.1)
        text(s, xx + 0.22, 4.52, GW - 0.4, 0.6, big, size=30, bold=True, color=INK)
        text(s, xx + 0.22, 5.22, GW - 0.4, 1.4, body, size=15, color=BODY)

    TX, TW = 8.75, 3.98
    label(s, TX, 1.65, TW, "Targets")
    for k, (big, body) in enumerate([("≥ 80%", "of Tier 1 flags are real barriers"),
                                     ("> 40%  ·  > 27%", "of obstacles · surface problems found, beating human labellers"),
                                     ("2 baselines", "beat zero-shot CLIP and a zero-shot vision-language model")]):
        yy = 2.02 + k * 1.64
        rect(s, TX, yy, TW, 1.52, fill=TEAL_SOFT if k == 0 else WHITE, line=None if k == 0 else RULE, radius=0.1)
        text(s, TX + 0.25, yy + 0.16, TW - 0.5, 0.55, big, size=26, bold=True, color=TEAL if k == 0 else INK)
        text(s, TX + 0.25, yy + 0.76, TW - 0.5, 0.7, body, size=14, color=BODY)

    # ---------------------------------------------------------------- 5. risk and help
    s = slide(prs, 5, "Risk  ·  help wanted", "Biggest risk: a model trained elsewhere may fail in Pittsburgh",
              NOTES[5])
    LW = 8.35

    def risk_label(x, y, w, n, t):
        text(s, x, y, w, 0.32, [[(f"{n}  ", {"color": TEAL}), (t, {"color": INK})]], size=15, bold=True)

    def fix(x, y, w, t):
        text(s, x, y, w, 0.3, [[("→  ", {"bold": True, "color": TEAL}), (t, {"color": TEAL})]], size=12.5)

    # risk 1: out-of-distribution cities (scope Q10)
    risk_label(0.6, 1.62, LW, 1, "Different city, different streets")
    PH = 1.9
    for x, img, tag in [(0.6, "rampnet_nyc", "Trained on: New York"), (3.08, "pgh_stairs_14480", "Pittsburgh: stairs"),
                        (5.08, "pgh_cobble_9626", "Pittsburgh: cobblestone")]:
        picture(s, IMG / f"{img}.jpg", x, 2.02, PH, PH)
        rect(s, x, 2.02 + PH - 0.3, PH, 0.3, fill=INK, alpha=0.72)
        text(s, x + 0.08, 2.02 + PH - 0.3, PH - 0.12, 0.3, tag, size=11, bold=True, color=WHITE, anchor="m")
    arrow(s, 2.58, 2.97, 2.98, 2.97, color=MUTED, lw=2.25)
    text(s, 7.2, 2.0, 1.75, 0.55, "1,128", size=28, bold=True, color=TEAL)
    text(s, 7.2, 2.55, 1.75, 0.5, "public stairways in Pittsburgh", size=12.5, color=BODY)
    text(s, 7.2, 3.15, 1.75, 0.75, "plus hills, brick and narrow sidewalks", size=12.5, color=MUTED)
    fix(0.6, 4.04, LW, "Test only on Pittsburgh, and check spots in person")

    CW2 = (LW - 0.25) / 2
    RY, RH = 4.5, 2.35
    # risk 2: imagery coverage of unaudited streets
    rect(s, 0.6, RY, CW2, RH, fill=SOFT, radius=0.1)
    risk_label(0.8, RY + 0.14, CW2 - 0.4, 2, "Photo coverage unknown")
    mw, mh = fit(B / "map_coverage.png", 1.55, 1.25)
    picture(s, B / "map_coverage.png", 0.8, RY + 0.56, mw, mh)
    chip(s, 0.8 + mw / 2 - 0.22, RY + 0.56 + mh / 2 - 0.22, 0.44, 0.44, "?", fill=INK, color=WHITE, size=18)
    text(s, 2.55, RY + 0.55, CW2 - 2.15, 0.45, "1,324 km", size=24, bold=True, color=INK)
    text(s, 2.55, RY + 1.0, CW2 - 2.15, 0.8, "of unaudited streets. How many have usable photos?", size=12.5, color=BODY)
    fix(0.8, RY + RH - 0.42, CW2 - 0.3, "Map photo coverage first")

    # risk 3: sparse labels across panoramas
    X3 = 0.6 + CW2 + 0.25
    rect(s, X3, RY, CW2, RH, fill=SOFT, radius=0.1)
    risk_label(X3 + 0.2, RY + 0.14, CW2 - 0.4, 3, "Sparse labels")
    for k in range(3):
        bx = X3 + 0.2 + k * (1.08 + 0.14)
        rect(s, bx, RY + 0.56, 1.08, 0.8, fill=WHITE, line=RULE, radius=0.06)
        rect(s, bx + 0.29, RY + 0.98, 0.5, 0.26, fill=SLATE, shape=MSO_SHAPE.TRAPEZOID)
        if k == 0:
            chip(s, bx + 0.08, RY + 0.63, 0.62, 0.24, "label", fill=TEAL, color=WHITE, size=9.5)
        else:
            text(s, bx + 0.1, RY + 0.63, 0.9, 0.24, "no label", size=9.5, color=MUTED)
    text(s, X3 + 0.2, RY + 1.47, CW2 - 0.4, 0.45, "One ramp, three photos, one label", size=12.5, color=BODY)
    fix(X3 + 0.2, RY + RH - 0.42, CW2 - 0.3, "RampNet labels every nearby photo")

    X, WR = 9.3, 3.43
    rect(s, X, 1.62, WR, 5.23, fill=TEAL_SOFT, radius=0.12)
    text(s, X + 0.25, 1.82, WR - 0.5, 0.45, "Where we'd love help", size=20, bold=True, color=TEAL)
    for k, (title, detail) in enumerate([("Imagery", "Photos of unaudited streets"),
                                         ("A DOMI planner", "Is our tier rule right? What list would you use?"),
                                         ("Disability advocates", "Review the tiers and the errors that matter most"),
                                         ("GPU hours", "To adapt the model to Pittsburgh")]):
        yy = 2.48 + k * 1.07
        chip(s, X + 0.25, yy + 0.03, 0.34, 0.34, str(k + 1), fill=TEAL, color=WHITE, size=12)
        text(s, X + 0.72, yy, WR - 0.95, 0.32, title, size=15, bold=True)
        text(s, X + 0.72, yy + 0.33, WR - 0.95, 0.62, detail, size=12.5, color=BODY)

    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
